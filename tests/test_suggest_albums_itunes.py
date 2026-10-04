"""suggest_albums_itunes: suggest only what MusicBrainz confirms, check what Grey typed,
and never touch Grey's own column."""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from musaeus.stages import album_fill as af

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "suggest_albums_itunes.py"


def _load():
    spec = importlib.util.spec_from_file_location("suggest_albums_itunes", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


def _it(track, album, date="1982-03-01", artist="Toto"):
    return {"artistName": artist, "trackName": track, "collectionName": album,
            "releaseDate": date + "T08:00:00Z", "trackCount": 10}  # fmt: skip


def test_suggests_only_confirmed_albums_and_checks_typed_ones(tmp_path, monkeypatch):
    mod = _load()
    wb = Workbook()
    ws = wb.active
    ws.append(["Album (type here)", "Artist", "Song"])
    ws.append([None, "Toto", "Africa"])  # iTunes proposes, MusicBrainz confirms
    ws.append([None, "Toto", "Hold The Line"])  # iTunes proposes, MusicBrainz does not
    ws.append([None, "Toto", "Rosanna (Live)"])  # a version title is never searched by text
    ws.append(["Toto IV", "Toto", "Rosanna"])  # Grey typed this one: check it, don't touch it
    w2 = wb.create_sheet("Check web-filled albums")
    w2.append(["Keep / change", "Artist", "Song", "Album filled"])
    w2.append([None, "Toto", "Africa", "Toto IV"])
    src = tmp_path / "w.xlsx"
    wb.save(src)

    monkeypatch.setattr(af.time, "sleep", lambda *_: None)
    monkeypatch.setattr(
        af, "itunes_search", lambda a, t: [_it(t, "Toto IV" if t == "Africa" else "Toto")]
    )
    monkeypatch.setattr(af, "mb_confirm", lambda a, al, t: (t in ("Africa", "Rosanna"), f"mb:{t}"))
    monkeypatch.setattr(mod, "policy", lambda *_a, **_k: __import__("contextlib").nullcontext())
    monkeypatch.setattr(sys, "argv", ["prog", str(src)])
    assert mod.main() == 0

    out = load_workbook(src)
    s1, s2 = out.worksheets[0], out["Check web-filled albums"]
    head = {c.value: c.column for c in s1[1]}
    assert s1.cell(2, head[mod.SUGGEST]).value == "Toto IV"
    assert s1.cell(3, head[mod.SUGGEST]).value is None
    assert "iTunes said 'Toto'" in s1.cell(3, head[mod.EVIDENCE]).value
    assert "version title" in s1.cell(4, head[mod.EVIDENCE]).value
    assert s1.cell(5, 1).value == "Toto IV", "Grey's own column is never touched"
    assert s1.cell(5, head[mod.CHECK]).value.startswith("CONFIRMED")
    h2 = {c.value: c.column for c in s2[1]}
    assert s2.cell(2, h2[mod.CHECK]).value.startswith("CONFIRMED")
    assert not src.with_suffix(".partial.xlsx").exists()


def test_discogs_comparison_adds_a_column_and_counts_who_confirms(tmp_path, monkeypatch, capsys):
    mod = _load()
    wb = Workbook()
    ws = wb.active
    ws.append(["Album (type here)", "Artist", "Song", mod.SUGGEST, mod.EVIDENCE])
    ws.append(
        [None, "Toto", "Africa", "Toto IV", "unambiguous; MusicBrainz: on the official album"]
    )
    ws.append(
        [
            None,
            "Toto",
            "Hold The Line",
            None,
            "iTunes said 'Toto' but MusicBrainz could not confirm",
        ]
    )
    ws.append(
        [None, "Toto", "Georgy Porgy", None, "no studio album for an exact artist+title match"]
    )
    src = tmp_path / "w.xlsx"
    wb.save(src)
    monkeypatch.setattr(af.time, "sleep", lambda *_: None)
    monkeypatch.setattr(
        af, "discogs_confirm", lambda a, al, t, k, s: (t == "Hold The Line", f"dg:{al}")
    )
    cfg = type("C", (), {"discogs_consumer_key": "k", "discogs_consumer_secret": "s"})()
    monkeypatch.setattr("musaeus.config.MusicConfig.from_env", staticmethod(lambda: cfg))
    monkeypatch.setattr(mod, "policy", lambda *_a, **_k: __import__("contextlib").nullcontext())
    monkeypatch.setattr(sys, "argv", ["prog", str(src), "--discogs"])
    assert mod.main() == 0
    s1 = load_workbook(src).worksheets[0]
    head = {c.value: c.column for c in s1[1]}
    assert s1.cell(2, head[mod.DISCOGS]).value.startswith("NOT CONFIRMED")
    assert s1.cell(3, head[mod.DISCOGS]).value.startswith("CONFIRMED")
    assert s1.cell(4, head[mod.DISCOGS]).value is None, "nothing proposed, nothing to check"
    out = capsys.readouterr().out
    assert "1  MusicBrainz only" in out and "1  Discogs only" in out
